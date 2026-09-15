"""일일 제조 수량표 엑셀 (차수별)."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from aggregator.batches import combine_batches, round_label, meat_group_subtotals
from aggregator.categories import DailyTotals
from aggregator.summary import production_rows

NAVY = "1F4E79"
GOLD = "FFF2CC"
WHITE = "FFFFFF"
GRAY = "F2F2F2"
GREEN = "E2EFDA"
BLUE = "D6EAF8"

thin = Border(
    left=Side(style="thin", color="7F8C8D"),
    right=Side(style="thin", color="7F8C8D"),
    top=Side(style="thin", color="7F8C8D"),
    bottom=Side(style="thin", color="7F8C8D"),
)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _apply(cell, *, font=None, fill=None):
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin


def build_report_workbook(
    totals: DailyTotals | None = None,
    report_date: date | datetime | None = None,
    batches: dict[int, DailyTotals] | None = None,
) -> Workbook:
    report_date = report_date or date.today()
    if isinstance(report_date, datetime):
        report_date = report_date.date()

    if batches:
        rounds = sorted(batches)
        rows = combine_batches(batches, only_ordered=False)
        qty_headers = [round_label(r) for r in rounds] + ["합계"]
    else:
        rounds = [1]
        assert totals is not None
        base_rows = production_rows(totals, only_ordered=False)
        rows = []
        for row in base_rows:
            rows.append(
                {
                    "상품": row["상품"],
                    "품목": row["품목"],
                    "용량": row["용량"],
                    "grams": row.get("grams") or 0,
                    "1차": row["총 개수"],
                    "합계": row["총 개수"],
                }
            )
        qty_headers = ["1차", "합계"]

    headers = ["상품", "품목", "용량", *qty_headers]
    last_col = len(headers)

    wb = Workbook()
    ws = wb.active
    ws.title = "차수별 총개수"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws["A1"]
    title.value = "당일 주문 용량별 제조 수량 (차수별)"
    _apply(title, font=Font(name="맑은 고딕", size=20, bold=True, color=WHITE), fill=_fill(NAVY))
    ws.row_dimensions[1].height = 34

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sub = ws["A2"]
    round_text = ", ".join(round_label(r) for r in rounds)
    sub.value = (
        f"{report_date.strftime('%Y년 %m월 %d일')}  /  {round_text}  "
        f"/  파일명 01·02·03 = 1차·2차·3차"
    )
    _apply(sub, font=Font(name="맑은 고딕", size=11, color=NAVY), fill=_fill(BLUE))
    ws.row_dimensions[2].height = 22

    for col, value in enumerate(headers, start=1):
        fill = GOLD if value == "합계" else NAVY
        color = NAVY if value == "합계" else WHITE
        _apply(
            ws.cell(4, col, value),
            font=Font(name="맑은 고딕", bold=True, color=color, size=12),
            fill=_fill(fill),
        )
    ws.row_dimensions[4].height = 24

    data_start = 5
    ordered_count = 0
    for idx, row in enumerate(rows):
        excel_row = data_start + idx
        total = row.get("합계") or 0
        if total:
            ordered_count += 1
        bg = GREEN if total else (GRAY if idx % 2 else WHITE)
        values = [row["상품"], row["품목"], row["용량"]] + [row.get(h, 0) or 0 for h in qty_headers]
        for col, value in enumerate(values, start=1):
            is_qty = col > 3
            is_sum = headers[col - 1] == "합계"
            font = Font(name="맑은 고딕", size=14, bold=True) if is_qty and total else Font(name="맑은 고딕", size=12)
            cell_fill = GOLD if is_sum else bg
            cell = ws.cell(excel_row, col, value)
            _apply(cell, font=font, fill=_fill(cell_fill))
        ws.row_dimensions[excel_row].height = 22

    group_rows = meat_group_subtotals(rows, rounds)
    group_title_row = data_start + len(rows) + 1
    ws.merge_cells(start_row=group_title_row, start_column=1, end_row=group_title_row, end_column=last_col)
    title_cell = ws.cell(group_title_row, 1)
    title_cell.value = "육회·육사시미 전용량 중량(kg)"
    _apply(
        title_cell,
        font=Font(name="맑은 고딕", bold=True, color=WHITE, size=12),
        fill=_fill(NAVY),
    )
    for col in range(2, last_col + 1):
        cell = ws.cell(group_title_row, col)
        cell.fill = _fill(NAVY)
        cell.border = thin

    group_header_row = group_title_row + 1
    for col, value in enumerate(["품목", "단위", ""] + qty_headers, start=1):
        fill = GOLD if value == "합계" else NAVY
        color = NAVY if value == "합계" else WHITE
        _apply(
            ws.cell(group_header_row, col, value),
            font=Font(name="맑은 고딕", bold=True, color=color, size=12),
            fill=_fill(fill),
        )

    for idx, grow in enumerate(group_rows):
        excel_row = group_header_row + 1 + idx
        values = [grow["품목"], grow.get("단위", "kg"), ""] + [grow.get(h, 0) for h in qty_headers]
        for col, value in enumerate(values, start=1):
            _apply(
                ws.cell(excel_row, col, value),
                font=Font(name="맑은 고딕", size=13, bold=True),
                fill=_fill("FCE4D6"),
            )

    sum_row = group_header_row + 1 + len(group_rows) + 1
    if last_col > 3:
        ws.merge_cells(start_row=sum_row, start_column=1, end_row=sum_row, end_column=3)
    sum_cell = ws.cell(sum_row, 1)
    sum_cell.value = "주문 있는 상품 수"
    _apply(sum_cell, font=Font(name="맑은 고딕", bold=True), fill=_fill(GOLD))
    for col in range(2, last_col):
        cell = ws.cell(sum_row, col)
        cell.fill = _fill(GOLD)
        cell.border = thin
    last_cell = ws.cell(sum_row, last_col)
    last_cell.value = ordered_count
    _apply(last_cell, font=Font(name="맑은 고딕", size=14, bold=True), fill=_fill(GOLD))

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12
    for i, _header in enumerate(qty_headers, start=4):
        ws.column_dimensions[get_column_letter(i)].width = 12

    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return wb


def report_bytes(totals=None, report_date=None, batches=None, **_kwargs) -> bytes:
    wb = build_report_workbook(totals, report_date, batches=batches)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
